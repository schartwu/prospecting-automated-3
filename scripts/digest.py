"""Build the morning digest email from Supabase and send it (or just write a preview).

Every weekday the same four lists: new D0 (bid now), B0 (prepare stock), 15 (reference prices) and P0 plans
scoring 75+. Only items first seen since the previous digest are included; results are on the web app.

Only Sales' provinces go in the email (TARGET_PROVINCES, a repository variable, else DEFAULT_PROVINCES).
Announcements whose province could not be read (mostly central offices in Bangkok/Nonthaburi) stay in;
those clearly in another province go to a separate Excel tab instead of being dropped.
"""
import argparse
import html
import os
import smtplib
import sys
from urllib.parse import quote
from datetime import datetime, timedelta
from email.message import EmailMessage
from email.utils import formataddr
from zoneinfo import ZoneInfo

import pandas as pd

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from egp import provinces  # noqa: E402
from egp.db import DB  # noqa: E402

BKK = ZoneInfo("Asia/Bangkok")
# KG Lead Loop web app (GitHub Pages, .github/workflows/pages.yml).
LEAD_APP_URL = os.environ.get("LEAD_APP_URL") or "https://schartwu.github.io/prospecting-automated-kg/"
ROOT = os.path.join(os.path.dirname(__file__), "..")
DEFAULT_PROVINCES = provinces.DEFAULT_TARGETS

SECTIONS = [
    # (types, heading, what Sales/Merchandise should do, min score or None for the default)
    # The same four every weekday; results, cancellations and summaries are on the web app.
    (["D0"], "🔥 ประกาศเชิญชวน — ยื่นราคาได้แล้ว", "เปิดเอกสาร ตรวจคุณสมบัติ และเตรียมยื่นราคา", None),
    (["B0"], "🛠 ร่าง TOR — เตรียมหาสินค้า (ประกาศจริงตามมาใน ~5 วัน)", "ให้ Merchandise ดู spec และเช็คสต็อก/ราคาทุน", None),
    (["15"], "💰 ประกาศราคากลาง", "รู้งบประมาณของโครงการ", None),
    (["P0"], "📅 แผนจัดซื้อใหม่ (คะแนน 75 ขึ้นไป)", "รู้ล่วงหน้าหลายเดือน วางแผนสินค้าและเข้าพบหน่วยงาน", 75),
]
TYPE_LABEL = {
    "P0": "แผนจัดซื้อ",
    "15": "ราคากลาง",
    "B0": "ร่าง TOR",
    "D0": "ประกาศเชิญชวน",
    "W0": "ประกาศผู้ชนะ",
    "D1": "ยกเลิกประกาศเชิญชวน",
    "W1": "ยกเลิกประกาศผู้ชนะ",
    "D2": "แก้ไขประกาศเชิญชวน",
    "W2": "แก้ไขประกาศผู้ชนะ",
}
TYPE_MEANING = {
    "P0": "แผนจัดซื้อจัดจ้างประจำปี — รู้ล่วงหน้าหลายเดือน ใช้วางแผนเข้าพบหน่วยงาน",
    "15": "ประกาศราคากลาง — บอกงบประมาณ ส่วนใหญ่ออกวันเดียวกับประกาศเชิญชวน",
    "B0": "ร่าง TOR / ร่างเอกสารประกวดราคา — ประกาศเชิญชวนจริงตามมาราว 5 วัน ให้เตรียมสินค้า",
    "D0": "ประกาศเชิญชวน — เปิดให้ยื่นราคาแล้ว ดูวันยื่นซอง",
    "W0": "ประกาศผู้ชนะ — ใครชนะ ใช้ดูคู่แข่งและหน่วยงานที่ซื้อของสำนักงานบ่อย",
    "D1": "ยกเลิกประกาศเชิญชวน",
    "W1": "ยกเลิกประกาศผู้ชนะ",
    "D2": "แก้ไข/เปลี่ยนแปลงประกาศเชิญชวน",
    "W2": "แก้ไข/เปลี่ยนแปลงประกาศผู้ชนะ",
}
SCORE_BANDS = [
    ("100-90", "สินค้าหลัก: วัสดุสำนักงาน กระดาษ หมึกพิมพ์ เครื่องเขียน"),
    ("89-75", "ซื้อ/เช่า คอมพิวเตอร์และไอที เฟอร์นิเจอร์ เครื่องใช้สำนักงาน"),
    ("74-60", "โครงการผสม มีสินค้าของเราเป็นส่วนสำคัญ"),
    ("59-40", "ก้ำกึ่ง: เครื่องใช้ไฟฟ้า เครื่องมือช่าง อุปกรณ์ความปลอดภัย"),
    ("39-0", "ไม่ใช่สินค้าที่ OfficeMate ขาย (ไม่แสดงในอีเมล)"),
]


def fmt_money(v):
    return f"{float(v):,.0f}" if v not in (None, "") and pd.notna(v) else "-"


def days_left(r, today):
    return (datetime.fromisoformat(r["bid_date"]).date() - today).days if r.get("bid_date") else None


def expected_label(r):
    """Plans carry the month the department expects to announce, e.g. '2026-10' -> 'ต.ค. 2569'."""
    if not r.get("expected_month"):
        return None
    year, month = (int(x) for x in r["expected_month"].split("-"))
    names = ["ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]
    return f"{names[month - 1]} {year + 543}"


def bid_cell(r, today):
    if not r.get("bid_date"):
        expected = expected_label(r)
        return f"<span style='color:#666;font-size:12px'>คาดว่าประกาศ</span><br>{expected}" if expected else "-"
    left = days_left(r, today)
    when = datetime.fromisoformat(r["bid_date"]).strftime("%d/%m")
    time_ = f" {r['bid_time']}" if r.get("bid_time") else ""
    color = "#c00" if left <= 3 else "#333"
    note = "ผ่านไปแล้ว" if left < 0 else ("วันนี้" if left == 0 else f"อีก {left} วัน")
    return f"{when}{time_}<div style='color:{color};font-size:12px'><b>{note}</b></div>"


EMAIL_ROW_LIMIT = 30  # per section; the attached Excel always has every row


def lead_key(r):
    """The key the Lead web app uses: the 11-digit project number, or 'P:' + plan number for plans."""
    return f"P:{r['plan_no']}" if r["announce_type"] == "P0" and r.get("plan_no") else r.get("project_no")


UPDATE, VIEW = "อัปเดตสถานะ", "ดูรายละเอียด"  # button text for DV/Sale, and for Monitors (read only)
# The Monitors' email goes to RECEIVER_EMAIL as one message, so it carries one shared read-only link: the
# lead_users row with this placeholder address (role viewer). Changing its token changes the link for all.
MONITOR_ACCOUNT = "monitor@kg-lead.local"
WEB_MORE = "ดูทั้งหมดในเว็บ KG Lead"


def update_button(url, label=UPDATE):
    return (f'<div style="margin-top:4px"><a href="{html.escape(url)}" style="display:inline-block;'
            'background:#0b6e69;color:#ffffff;text-decoration:none;font-size:12px;padding:3px 10px;'
            f'border-radius:5px">{label}</a></div>')


def section_html(heading, hint, rows, today, app_link=None, label=UPDATE, excel=True):
    """app_link(key) -> the reader's link to that lead in the web app, or None (then the Excel file is attached)."""
    out = [f'<h3 style="margin:24px 0 4px">{html.escape(heading)} ({len(rows)})</h3>',
           f'<div style="color:#666;font-size:13px;margin-bottom:8px">{html.escape(hint)}</div>']
    shown, rows = rows, rows[:EMAIL_ROW_LIMIT]
    if not rows:
        out.append('<div style="color:#999;font-size:13px">ไม่มีรายการใหม่</div>')
        return "\n".join(out)
    out.append('<table cellpadding="6" style="border-collapse:collapse;font-size:13px;width:100%">'
               '<tr style="background:#f2f2f2;text-align:left"><th>คะแนน</th><th>โครงการ</th><th>หน่วยงาน</th>'
               '<th style="white-space:nowrap">งบ (บาท)</th><th style="white-space:nowrap">วันยื่นซอง</th>'
               '<th style="white-space:nowrap">วันที่ประกาศ</th></tr>')
    for r in rows:
        flags = ", ".join(TYPE_LABEL.get(s, s) for s in (r.get("stages") or []) if s != r["announce_type"])
        badges = ""
        if r.get("cancelled"):
            badges += ' <b style="color:#c00">[ยกเลิกแล้ว]</b>'
        if r.get("amended"):
            badges += ' <b style="color:#d97706">[มีการแก้ไขประกาศ]</b>'
        sub = " · ".join(filter(None, [r.get("category"), r.get("reason"),
                                       f"มีประกาศ: {flags}" if flags else None]))
        if app_link and lead_key(r):
            badges += update_button(app_link(lead_key(r)), label)
        out.append(
            "<tr style='border-top:1px solid #eee;vertical-align:top'>"
            f"<td><b>{r['score']}</b></td>"
            f"<td><a href=\"{html.escape(r['link'])}\">{html.escape(r['title'][:220])}</a>{badges}"
            f"<div style='color:#888;font-size:12px'>{html.escape(sub)}</div></td>"
            f"<td>{html.escape(r.get('dept_name') or '')}"
            f"<div style='color:#888;font-size:12px'>{html.escape(r.get('province') or 'ไม่ระบุจังหวัด')}</div></td>"
            f"<td style='text-align:right;white-space:nowrap'>{fmt_money(r.get('budget_thb'))}</td>"
            f"<td style='white-space:nowrap'>{bid_cell(r, today)}</td>"
            f"<td style='white-space:nowrap'>{r.get('pub_date') or ''}</td></tr>")
    out.append("</table>")
    if len(shown) > len(rows):
        out.append(f'<div style="color:#666;font-size:12px">แสดง {len(rows)} อันดับแรก '
                   f'จาก {len(shown)} รายการ · {"ดูทั้งหมดในไฟล์ Excel แนบ" if excel else WEB_MORE}</div>')
    return "\n".join(out)


def build(db: DB, since: datetime, now: datetime, min_score: int, targets: set[str],
          app_link=None, board_link=None, label=UPDATE):
    """Returns (email html, rows in the email, rows clearly in other provinces).
    With app_link/board_link the lists live on the web app and no Excel is attached."""
    excel = board_link is None
    today = now.date()
    parts, all_rows, elsewhere = [], [], []
    for types, heading, hint, sec_min in SECTIONS:
        rows = db.rpc("digest_items", {"p_since": since.isoformat(), "p_types": types,
                                       "p_min_score": max(min_score, sec_min or 0)})
        # Keep Sales' provinces and the ones nobody could place; set the rest aside for the Excel tab.
        elsewhere += [{**r, "section": heading, "days_left": days_left(r, today)} for r in rows
                      if r.get("province") and r["province"] not in targets]
        rows = [r for r in rows if not r.get("province") or r["province"] in targets]
        if types == ["15"]:
            # Reference prices for projects that also got a D0 are already shown under D0.
            rows = [r for r in rows if "D0" not in (r.get("stages") or [])]
        if types == ["D0"]:
            # Bids closing soonest first; unknown or past deadlines after them.
            def urgency(r):
                left = days_left(r, today)
                return (left is None or left < 0, left if left is not None else 999, -r["score"])
            rows.sort(key=urgency)
        parts.append(section_html(heading, hint, rows, today, app_link=app_link, label=label, excel=excel))
        all_rows += [{**r, "section": heading, "days_left": days_left(r, today)} for r in rows]

    total = len(all_rows)
    head = (f'<div style="font-family:Tahoma,sans-serif;max-width:980px">'
            f'<h2 style="margin:0">e-GP Alert — {now:%d/%m/%Y}</h2>'
            f'<div style="color:#666">ประกาศใหม่ที่เกี่ยวกับ OfficeMate ตั้งแต่ {since.astimezone(BKK):%d/%m %H:%M} น. '
            f'รวม {total} รายการ (คะแนน ≥ {min_score})<br>เฉพาะจังหวัด: {html.escape(", ".join(sorted(targets)))} '
            f'และประกาศที่ไม่ระบุจังหวัด (ส่วนใหญ่เป็นส่วนกลาง)'
            + (f' · จังหวัดอื่น {len(elsewhere)} รายการ'
               + (' ดูในเว็บ KG Lead เลือกพื้นที่ "จังหวัดอื่น"' if board_link
                  else ' อยู่ในแท็บ "รายการ (จังหวัดอื่น)" ของไฟล์ Excel') if elsewhere else '')
            + '</div>'
            + (f'<div style="margin-top:8px"><a href="{html.escape(board_link)}" style="color:#0b6e69;font-weight:bold">'
               'เปิด KG Lead: รายการทั้งหมด ผลลัพธ์ และสถานะ (ย้อนดูได้ทุกวัน) →</a></div>' if board_link else ''))
    bands = " · ".join(f"{b} {d}" for b, d in SCORE_BANDS[:4])
    foot = (f'<p style="color:#999;font-size:12px;margin-top:24px">คะแนน = ความตรงกับสินค้า OfficeMate: '
            f'{html.escape(bands)}</p></div>')
    return head + "\n".join(parts) + foot, all_rows, elsewhere


def write_xlsx(rows, path, elsewhere=()):
    cols = {"section": "หมวดอีเมล", "score": "คะแนน", "type_label": "ประเภทประกาศ", "title": "ชื่อโครงการ",
            "dept_name": "หน่วยงาน", "province": "จังหวัด", "category": "หมวดสินค้า", "reason": "เหตุผล",
            "budget_thb": "งบประมาณ (บาท)", "winners": "ผู้ชนะ", "winner_price": "ราคาที่ชนะ (บาท)",
            "bid_date": "วันยื่นซอง", "bid_time": "เวลายื่น", "days_left": "เหลือ (วัน)",
            "expected": "คาดว่าจะประกาศ",
            "pub_date": "วันที่ประกาศ", "method": "วิธีจัดหา", "project_no": "เลขโครงการ", "plan_no": "เลขแผน",
            "cancelled": "ยกเลิกแล้ว", "amended": "มีการแก้ไขประกาศ", "link": "ลิงก์"}
    def frame(rs):
        data = [{**r, "type_label": f"{TYPE_LABEL.get(r['announce_type'], r['announce_type'])} ({r['announce_type']})",
                 "expected": expected_label(r)} for r in rs]
        return pd.DataFrame(data, columns=list(cols)).rename(columns=cols)

    columns_help = [
        ("จังหวัด", "อ่านจากชื่อโครงการและเอกสารประกาศ (ว่าง = เอกสารไม่ระบุจังหวัด ส่วนใหญ่เป็นหน่วยงานส่วนกลาง)"),
        ("งบประมาณ (บาท)", "ราคากลางจากเอกสารประกาศ หรืองบประมาณโครงการจากหน้าแผนจัดซื้อ (ว่าง = เอกสารต้องล็อกอิน e-GP หรือไม่ระบุ "
                           "ประกาศผู้ชนะแบบเฉพาะเจาะจงไม่มีประกาศราคากลางก่อนหน้า จึงมักว่าง ให้ดูราคาที่ชนะแทน)"),
        ("ผู้ชนะ / ราคาที่ชนะ (บาท)", "ชื่อผู้ชนะและราคาที่เสนอจากประกาศผู้ชนะ (W0) ถ้ามีหลายรายการ ราคาเป็นยอดรวม"),
        ("วันยื่นซอง / เวลายื่น", "วันและช่วงเวลาเสนอราคาในระบบ e-GP (มีเฉพาะประกาศเชิญชวน)"),
        ("คาดว่าจะประกาศ", "เดือนที่หน่วยงานคาดว่าจะประกาศจัดซื้อ (มีเฉพาะแผนจัดซื้อ P0)"),
        ("เหลือ (วัน)", "จำนวนวันจากวันที่ส่งอีเมลถึงวันยื่นซอง"),
        ("ยกเลิกแล้ว", "โครงการเดียวกันมีประกาศยกเลิกประกาศเชิญชวน (D1) ตามมา — ไม่ต้องตามต่อ"),
        ("มีการแก้ไขประกาศ", "โครงการเดียวกันมีประกาศแก้ไขประกาศเชิญชวน (D2) — เปิดดูรายละเอียดใหม่"),
        ("เลขโครงการ", "เลข 11 หลักของ e-GP ใช้ค้นหาโครงการในเว็บ e-GP"),
    ]
    legend = ([{"หัวข้อ": "ประเภทประกาศ", "ค่า": f"{TYPE_LABEL[k]} ({k})", "ความหมาย": v}
               for k, v in TYPE_MEANING.items()]
              + [{"หัวข้อ": "คะแนน", "ค่า": b, "ความหมาย": d} for b, d in SCORE_BANDS]
              + [{"หัวข้อ": "คอลัมน์", "ค่า": k, "ความหมาย": v} for k, v in columns_help])
    sheets = [("รายการ (กทม. และปริมณฑล)", rows)] + ([("รายการ (จังหวัดอื่น)", elsewhere)] if elsewhere else [])
    with pd.ExcelWriter(path) as xw:
        for name, rs in sheets:
            frame(rs).to_excel(xw, sheet_name=name, index=False)
        pd.DataFrame(legend).to_excel(xw, sheet_name="คำอธิบาย", index=False)
        for name, _ in sheets:
            style_sheet(xw.sheets[name], number_cols={"งบประมาณ (บาท)", "ราคาที่ชนะ (บาท)"},
                        wrap_cols={"ชื่อโครงการ": 60, "เหตุผล": 35, "หน่วยงาน": 28, "ผู้ชนะ": 30,
                                   "หมวดอีเมล": 24, "หมวดสินค้า": 22, "วิธีจัดหา": 24})
        style_sheet(xw.sheets["คำอธิบาย"], wrap_cols={"ความหมาย": 80})


# Column widths in Excel character units. Long text columns (wrap_cols: name -> width) get a fixed
# width and wrap instead of stretching the sheet; the rest are sized to their content within bounds.
MIN_WIDTH, MAX_WIDTH = 8, 45


def style_sheet(ws, number_cols=frozenset(), wrap_cols=None):
    wrap_cols = wrap_cols or {}
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header = {cell.value: cell.column for cell in ws[1]}
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F2F2F2")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for name, col in header.items():
        letter = get_column_letter(col)
        cells = [c for (c,) in ws.iter_rows(min_row=2, min_col=col, max_col=col)]
        if name in number_cols:
            for c in cells:
                c.number_format = "#,##0"  # 4,992,128 — no decimals
        if name in wrap_cols:
            width = wrap_cols[name]
            for c in cells:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        else:
            longest = max([len(str(name))] + [len(f"{c.value:,.0f}") if name in number_cols and isinstance(
                c.value, (int, float)) else len(str(c.value)) for c in cells if c.value is not None])
            width = min(MAX_WIDTH, max(MIN_WIDTH, longest + 2))
            for c in cells:
                c.alignment = Alignment(vertical="top")
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


SENDER_NAME = "KG Lead (e-GP Alert)"  # shown instead of the bare sender address


def send(subject, html_body, attachment, recipients, sender_name=None):
    """sender_name is the name shown instead of the bare address (the address itself stays visible)."""
    msg = EmailMessage()
    msg["From"] = formataddr((sender_name, os.environ["SENDER_EMAIL"])) if sender_name else os.environ["SENDER_EMAIL"]
    msg["To"] = ", ".join(recipients)
    msg["Subject"] = subject
    msg.set_content("กรุณาเปิดอีเมลนี้ในโหมด HTML")
    msg.add_alternative(html_body, subtype="html")
    if attachment:  # the Lead Loop trial copy has none: its lists are on the web app
        with open(attachment, "rb") as f:
            msg.add_attachment(f.read(), maintype="application",
                               subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               filename=os.path.basename(attachment))
    with smtplib.SMTP_SSL("smtp.gmail.com", 465) as smtp:
        smtp.login(os.environ["SENDER_EMAIL"], os.environ["SENDER_PASSWORD"])
        smtp.send_message(msg)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--send", action="store_true", help="actually email it (default: preview files only)")
    ap.add_argument("--to", default="", help="override recipients (comma separated) for tests")
    ap.add_argument("--monitor", action="store_true", help="with --to: send the Monitors' version instead")
    ap.add_argument("--since-days", type=float, default=0, help="ignore the last digest; look back N days")
    ap.add_argument("--min-score", type=int, default=50)
    args = ap.parse_args()

    db = DB()
    now = datetime.now(BKK)
    last = db.select("digests", kind="eq.daily", order="sent_at.desc", limit=1)
    if args.since_days:
        since = now - timedelta(days=args.since_days)
    elif last:
        since = datetime.fromisoformat(last[0]["sent_at"])
    else:
        since = now - timedelta(days=1)

    targets = provinces.parse_targets(os.environ.get("TARGET_PROVINCES") or DEFAULT_PROVINCES)
    out = os.path.join(ROOT, "out")
    os.makedirs(out, exist_ok=True)

    def make(suffix: str = "", token: str | None = None, label: str = UPDATE):
        """With the reader's token: a button on every lead (label) and a link to the web app, no Excel."""
        app_link = (lambda key: f"{LEAD_APP_URL}#t={token}&lead={quote(key)}") if token else None
        body, rows, elsewhere = build(db, since, now, args.min_score, targets, app_link=app_link, label=label,
                                      board_link=f"{LEAD_APP_URL}#t={token}" if token else None)
        xlsx = os.path.join(out, f"eGP_Alert_{now:%Y%m%d}{suffix}.xlsx")
        write_xlsx(rows, xlsx, elsewhere)
        if not token:  # out/ is uploaded as a run artifact of this public repository; personal links stay out
            with open(os.path.join(out, f"digest_preview{suffix}.html"), "w", encoding="utf-8") as f:
                f.write(body)
        print(f"digest{suffix} since {since:%Y-%m-%d %H:%M}: {len(rows)} items "
              f"(+{len(elsewhere)} in other provinces) -> {out}")
        return body, rows, xlsx

    def subject(rows, prefix=""):
        n_bid = sum(r["announce_type"] == "D0" for r in rows)
        n_prep = sum(r["announce_type"] == "B0" for r in rows)
        n_plan = sum(r["announce_type"] == "P0" for r in rows)
        return prefix + (f"🚨 e-GP Alert {now:%d/%m/%Y}: ยื่นราคาได้ {n_bid} · เตรียมของ {n_prep} · "
                         f"{f'แผนใหม่ {n_plan} · ' if n_plan else ''}ใหม่รวม {len(rows)} โครงการ" if rows
                         else f"ℹ️ e-GP Alert {now:%d/%m/%Y}: ไม่มีโครงการใหม่")

    def emails(value):
        return [e.strip() for e in (value or "").split(",") if e.strip()]

    def token_of(email):
        """The person's own web-app link token, if they are a Lead Loop user. Each of them gets their own
        copy of the email, since the links inside identify the reader."""
        found = db.select("lead_users", select="token", email=f"eq.{email.lower()}", active="eq.true")
        return found[0]["token"] if found else None

    def send_personal(to, prefix, suffix):
        for i, email in enumerate(to):
            token = token_of(email)
            body, rows, xlsx = make(suffix=f"{suffix}{i}", token=token)
            if token is None:
                print(f"::warning::{email} is not a Lead Loop user; sent without web links")
            send(subject(rows, prefix), body, None if token else xlsx, [email], sender_name=SENDER_NAME)
        print(f"{prefix.strip() or 'personal copies'} sent to {len(to)} recipient(s)")

    def editors():
        """Everyone who updates leads (DV and Sales), from the KG Lead user list."""
        rows = db.select("lead_users", select="email", role="in.(dv,sale)", active="eq.true", order="id")
        return [r["email"] for r in rows]

    # LOOP_RECIPIENTS: "all" = every active DV/Sale gets their own copy; otherwise only the addresses listed
    # (while trying things out).
    loop_value = (os.environ.get("LOOP_RECIPIENTS") or "").strip()
    loop_to = editors() if loop_value.lower() == "all" else emails(loop_value)
    if not args.send:
        make()  # previews show everything, trial parts included (no personal links)
        return

    if args.to:  # a test send, and the "since" window of the real digest stays put
        if args.monitor:
            token = token_of(MONITOR_ACCOUNT)
            body, rows, xlsx = make(suffix="_test_monitor", token=token, label=VIEW)
            send(subject(rows), body, None if token else xlsx, emails(args.to),
                 sender_name=SENDER_NAME)
            print("Monitor test copy sent")
        else:
            send_personal(emails(args.to), "", "_test")
        return

    # DIGEST_RECIPIENTS (a repository variable) overrides the team list, e.g. for a trial period.
    recipients = emails(os.environ.get("DIGEST_RECIPIENTS") or os.environ["RECEIVER_EMAIL"])
    # The Monitors read only: one message with the shared read-only link; the lists are on the web app.
    # Without that account the old team email (Excel attached, no links) goes out instead.
    monitor_token = token_of(MONITOR_ACCOUNT)
    if monitor_token is None:
        print(f"::warning::{MONITOR_ACCOUNT} is not in lead_users; team email sent with Excel and no links")
    body, rows, xlsx = make(suffix="_team", token=monitor_token, label=VIEW)
    send(subject(rows), body, None if monitor_token else xlsx, recipients, sender_name=SENDER_NAME)
    db.insert("digests", {"kind": "daily", "since": since.isoformat(), "n_items": len(rows),
                          "recipients": ",".join(recipients)}, returning=False)
    print(f"sent to {len(recipients)} recipient(s)")
    if loop_to:  # every day: each DV/Sale their own copy with the "อัปเดตสถานะ" buttons
        send_personal(loop_to, "", "_personal")


if __name__ == "__main__":
    main()
